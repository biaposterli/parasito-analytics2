"""Configuração da análise — parasitos, classificação, métodos e nº de amostras.

Versão híbrida:
  * se a planilha enviada tiver as abas Config_Parasitos / Config_Metodos /
    Config_Amostras, elas pré-preenchem a configuração;
  * se não tiver, a configuração é montada a partir do que foi detectado na
    própria planilha (colunas de método e parasitos encontrados nos resultados);
  * em ambos os casos o usuário revisa/edita no site, e pode baixar a
    configuração pronta (sozinha ou já anexada à planilha original) para não
    precisar configurar de novo.

As tabelas trocadas com o site (st.data_editor) e com a planilha usam os mesmos
cabeçalhos, definidos abaixo.
"""
from __future__ import annotations

import io
import unicodedata

import pandas as pd

from analysis_engine import (
    CATEGORIAS_VALIDAS,
    DOMINIO_FECAL,
    DOMINIO_LAMINA,
    METHOD_CATALOG,
    PARASITE_MAP,
    CRITERIOS_PADRAO,
    TODAS,
    AnalysisConfig,
    _parse_result_cell_raw,
    categoria_de,
    max_coletas,
)

from estilo_planilha import estilizar_config

SHEET_PARASITOS = "Config_Parasitos"
SHEET_METODOS = "Config_Metodos"
SHEET_AMOSTRAS = "Config_Amostras"
SHEET_CRITERIOS = "Config_Criterios"
CONFIG_SHEETS = (SHEET_PARASITOS, SHEET_METODOS, SHEET_AMOSTRAS, SHEET_CRITERIOS)

AMOSTRA_FEZES = "Fezes"
AMOSTRA_LAMINA = "Lâmina (Graham)"
AMOSTRAS_VALIDAS = (AMOSTRA_FEZES, AMOSTRA_LAMINA)

PAR_COLS = ["Incluir", "Parasito", "Classificação", "Agrupar como", "Encontrado em", "Ocorrências"]
MET_COLS = ["Incluir", "Coluna", "Método", "Amostra"]
PARAM_N_AMOSTRAS = "Quantidade de amostras por paciente"

# Aba Config_Criterios: (chave interna, rótulo na planilha, explicação)
CRITERIOS_ROTULOS = [
    ("min_fezes", "Mínimo de coletas com fezes entregues",
     "0 = sem exigência; um número (1, 2...) ou 'Todas' (todas as coletas P1..Pn consideradas)."),
    ("min_lamina", "Mínimo de coletas com lâmina entregue",
     "0 = sem exigência; um número (1, 2...) ou 'Todas'."),
    ("excluir_inconclusivos", "Excluir paciente com amostra sem resultado",
     "Sim = exclui quem teve 'Amostra insuficiente' ou célula vazia em algum método incluído."),
    ("exigir_todos_metodos", "Exigir todos os métodos realizados",
     "Sim = exclui quem teve algum método incluído marcado 'Não realizado'."),
]


def _criterio_valor_planilha(chave, valor):
    if chave in ("min_fezes", "min_lamina"):
        return "Todas" if valor == TODAS else int(valor or 0)
    return "Sim" if valor else "Não"


def _criterio_da_planilha(chave, valor):
    if chave in ("min_fezes", "min_lamina"):
        k = _key(valor)
        if k.startswith("tod"):
            return TODAS
        try:
            return max(0, int(float(valor)))
        except (TypeError, ValueError):
            return 0
    return _as_bool(valor, default=False)


# ----------------------------------------------------------------------
# utilidades de leitura
# ----------------------------------------------------------------------
def _key(x) -> str:
    """minúsculas, sem acento, sem espaços nas pontas — para comparar cabeçalhos."""
    t = "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x)
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    return " ".join(t.strip().lower().split())


def _as_bool(x, default=True) -> bool:
    if isinstance(x, bool):
        return x
    k = _key(x)
    if k == "" or k == "nan":
        return default
    if k in {"nao", "n", "no", "false", "falso", "0", "excluir"}:
        return False
    return True


def _as_categoria(x) -> str:
    k = _key(x)
    if k.startswith("pat"):
        return "Patogênico"
    if k.startswith("com"):
        return "Comensal"
    return "Não classificado"


def _as_amostra(x) -> str:
    k = _key(x)
    if any(w in k for w in ("lam", "graham", "swab", "fita", "perianal")):
        return AMOSTRA_LAMINA
    return AMOSTRA_FEZES


def _dominio(amostra: str) -> str:
    return DOMINIO_LAMINA if amostra == AMOSTRA_LAMINA else DOMINIO_FECAL


def _amostra(dominio: str) -> str:
    return AMOSTRA_LAMINA if dominio == DOMINIO_LAMINA else AMOSTRA_FEZES


def _texto(x) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return " ".join(str(x).split())


def _renomear_colunas(df: pd.DataFrame, esperadas: list[str]) -> pd.DataFrame:
    by_key = {_key(c): c for c in esperadas}
    return df.rename(columns={c: by_key[_key(c)] for c in df.columns if _key(c) in by_key})


# ----------------------------------------------------------------------
# detecção a partir da planilha de resultados
# ----------------------------------------------------------------------
_CATALOG_BY_COL = {col: (nome, dom) for col, nome, _, dom in METHOD_CATALOG}


def _nome_padrao_metodo(col: str) -> str:
    if col in _CATALOG_BY_COL:
        return _CATALOG_BY_COL[col][0]
    base = col[len("metodo_"):] if col.startswith("metodo_") else col
    return base.replace("_", " ").strip().title() or col


def _dominio_padrao_metodo(col: str) -> str:
    if col in _CATALOG_BY_COL:
        return _CATALOG_BY_COL[col][1]
    k = _key(col)
    return DOMINIO_LAMINA if any(w in k for w in ("graham", "swab", "lamina", "fita")) else DOMINIO_FECAL


def detectar_colunas_metodo(df: pd.DataFrame) -> list[str]:
    """Colunas de método presentes: as do catálogo + qualquer outra 'metodo_*'."""
    cat = [col for col, _, _, _ in METHOD_CATALOG if col in df.columns]
    extras = [c for c in df.columns if str(c).startswith("metodo_") and c not in cat]
    return cat + extras


def tabela_metodos_detectados(df: pd.DataFrame) -> pd.DataFrame:
    rows = [
        {"Incluir": True, "Coluna": col, "Método": _nome_padrao_metodo(col),
         "Amostra": _amostra(_dominio_padrao_metodo(col))}
        for col in detectar_colunas_metodo(df)
    ]
    return pd.DataFrame(rows, columns=MET_COLS)


def contar_parasitos(df: pd.DataFrame, metodos: pd.DataFrame) -> pd.DataFrame:
    """Parasitos encontrados nas colunas de método incluídas: nº de células com o
    parasito e em que tipo de amostra (fezes / lâmina) ele apareceu."""
    cont: dict[str, int] = {}
    onde: dict[str, set] = {}
    incl = metodos[metodos["Incluir"].apply(_as_bool)] if not metodos.empty else metodos
    for _, m in incl.iterrows():
        col = m["Coluna"]
        if col not in df.columns:
            continue
        dominio = _dominio(m["Amostra"])
        for raw in df[col]:
            _, especies, positivo = _parse_result_cell_raw(raw)
            if not positivo:
                continue
            for e in especies:
                cont[e] = cont.get(e, 0) + 1
                onde.setdefault(e, set()).add(dominio)
    rows = []
    for e in sorted(cont, key=lambda k: (-cont[k], k)):
        d = onde.get(e, set())
        enc = " + ".join(x for x in (AMOSTRA_FEZES if DOMINIO_FECAL in d else "",
                                     "Lâmina" if DOMINIO_LAMINA in d else "") if x)
        rows.append({"Parasito": e, "Encontrado em": enc, "Ocorrências": cont[e]})
    return pd.DataFrame(rows, columns=["Parasito", "Encontrado em", "Ocorrências"])


def tabela_parasitos_padrao(contagem: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in contagem.iterrows():
        rows.append({
            "Incluir": True, "Parasito": r["Parasito"],
            "Classificação": categoria_de(r["Parasito"]),
            "Agrupar como": "", "Encontrado em": r["Encontrado em"],
            "Ocorrências": int(r["Ocorrências"]),
        })
    return pd.DataFrame(rows, columns=PAR_COLS)


# ----------------------------------------------------------------------
# leitura das abas Config_* da planilha
# ----------------------------------------------------------------------
def ler_config_da_planilha(xls: pd.ExcelFile) -> dict:
    """Lê as abas de configuração, se existirem. Devolve um dict com as chaves
    'parasitos', 'metodos', 'n_amostras' (cada uma só se a aba existir)."""
    by_key = {_key(s): s for s in xls.sheet_names}
    out: dict = {}

    s = by_key.get(_key(SHEET_PARASITOS))
    if s:
        d = _renomear_colunas(pd.read_excel(xls, sheet_name=s), PAR_COLS)
        if "Parasito" in d.columns:
            d = d[d["Parasito"].apply(_texto) != ""].copy()
            d["Parasito"] = d["Parasito"].apply(_texto)
            d["Incluir"] = d["Incluir"].apply(_as_bool) if "Incluir" in d.columns else True
            d["Classificação"] = d["Classificação"].apply(_as_categoria) if "Classificação" in d.columns \
                else d["Parasito"].apply(categoria_de)
            d["Agrupar como"] = d["Agrupar como"].apply(_texto) if "Agrupar como" in d.columns else ""
            out["parasitos"] = d[["Incluir", "Parasito", "Classificação", "Agrupar como"]].reset_index(drop=True)

    s = by_key.get(_key(SHEET_METODOS))
    if s:
        d = _renomear_colunas(pd.read_excel(xls, sheet_name=s), MET_COLS)
        if "Coluna" in d.columns:
            d = d[d["Coluna"].apply(_texto) != ""].copy()
            d["Coluna"] = d["Coluna"].apply(lambda c: _texto(c).lower().replace(" ", "_"))
            d["Incluir"] = d["Incluir"].apply(_as_bool) if "Incluir" in d.columns else True
            d["Método"] = [(_texto(n) or _nome_padrao_metodo(c)) for n, c in
                           zip(d["Método"] if "Método" in d.columns else [""] * len(d), d["Coluna"])]
            d["Amostra"] = [(_as_amostra(a) if _texto(a) else _amostra(_dominio_padrao_metodo(c))) for a, c in
                            zip(d["Amostra"] if "Amostra" in d.columns else [""] * len(d), d["Coluna"])]
            out["metodos"] = d[MET_COLS].reset_index(drop=True)

    s = by_key.get(_key(SHEET_AMOSTRAS))
    if s:
        d = pd.read_excel(xls, sheet_name=s, header=None)
        for _, r in d.iterrows():
            vals = [v for v in r.tolist() if not (isinstance(v, float) and pd.isna(v))]
            if len(vals) >= 2 and "amostra" in _key(vals[0]):
                try:
                    out["n_amostras"] = int(float(vals[1]))
                except (TypeError, ValueError):
                    pass
                break

    s = by_key.get(_key(SHEET_CRITERIOS))
    if s:
        d = pd.read_excel(xls, sheet_name=s, header=None)
        por_rotulo = {_key(r): ch for ch, r, _ in CRITERIOS_ROTULOS}
        crit = dict(CRITERIOS_PADRAO)
        achou = False
        for _, r in d.iterrows():
            vals = r.tolist()
            if len(vals) >= 2 and _key(vals[0]) in por_rotulo:
                ch = por_rotulo[_key(vals[0])]
                crit[ch] = _criterio_da_planilha(ch, vals[1])
                achou = True
        if achou:
            out["criterios"] = crit
    return out


# ----------------------------------------------------------------------
# combinação: configuração da planilha + o que foi detectado
# ----------------------------------------------------------------------
def montar_tabelas_iniciais(df: pd.DataFrame, cfg_planilha: dict) -> tuple[pd.DataFrame, pd.DataFrame, int, int, list[str]]:
    """Devolve (tabela_parasitos, tabela_metodos, n_amostras, max_observado, avisos).

    Regras de combinação quando existe configuração na planilha:
      * métodos listados na configuração mas sem coluna nos dados são avisados e
        descartados; colunas metodo_* nos dados que a configuração não cita
        entram como incluídas (e são avisadas);
      * parasitos da configuração que não apareceram nos dados continuam na
        lista (0 ocorrências); parasitos encontrados nos dados e ausentes da
        configuração entram com a classificação padrão (e são avisados).
    """
    avisos: list[str] = []

    detectados = tabela_metodos_detectados(df)
    if "metodos" in cfg_planilha:
        met = cfg_planilha["metodos"]
        faltando = [c for c in met["Coluna"] if c not in df.columns]
        if faltando:
            avisos.append("Métodos da configuração sem coluna nos dados (ignorados): " + ", ".join(faltando) + ".")
        met = met[met["Coluna"].isin(df.columns)]
        novos = detectados[~detectados["Coluna"].isin(met["Coluna"])]
        if not novos.empty:
            avisos.append("Colunas de método que não estavam na configuração (incluídas): "
                          + ", ".join(novos["Coluna"]) + ".")
        metodos = pd.concat([met, novos], ignore_index=True)[MET_COLS]
    else:
        metodos = detectados

    contagem = contar_parasitos(df, metodos)
    if "parasitos" in cfg_planilha:
        par = cfg_planilha["parasitos"].drop_duplicates("Parasito").copy()
        info = contagem.set_index("Parasito")
        par["Encontrado em"] = [info["Encontrado em"].get(p, "—") if p in info.index else "—" for p in par["Parasito"]]
        par["Ocorrências"] = [int(info["Ocorrências"].get(p, 0)) if p in info.index else 0 for p in par["Parasito"]]
        novos = tabela_parasitos_padrao(contagem[~contagem["Parasito"].isin(par["Parasito"])])
        if not novos.empty:
            avisos.append("Parasitos encontrados nos dados que não estavam na configuração (incluídos com a "
                          "classificação padrão — confira): " + ", ".join(novos["Parasito"]) + ".")
        parasitos = pd.concat([par, novos], ignore_index=True)[PAR_COLS]
    else:
        parasitos = tabela_parasitos_padrao(contagem)

    max_obs = max_coletas(df)
    n = cfg_planilha.get("n_amostras") or max_obs or 1
    if max_obs and n > max_obs:
        avisos.append(f"A configuração pede {n} amostras por paciente, mas a planilha só vai até P{max_obs}.")
    return parasitos.reset_index(drop=True), metodos.reset_index(drop=True), int(n), int(max_obs), avisos


def atualizar_contagens(parasitos: pd.DataFrame, df: pd.DataFrame, metodos: pd.DataFrame) -> pd.DataFrame:
    """Recalcula 'Encontrado em'/'Ocorrências' depois que o usuário muda os
    métodos (incluir/excluir, fezes/lâmina) e acrescenta parasitos que passaram
    a aparecer, com a classificação padrão. Mantém as escolhas já feitas."""
    contagem = contar_parasitos(df, metodos)
    info = contagem.set_index("Parasito")
    par = parasitos.copy()
    par["Encontrado em"] = [info["Encontrado em"].get(p, "—") if p in info.index else "—" for p in par["Parasito"]]
    par["Ocorrências"] = [int(info["Ocorrências"].get(p, 0)) if p in info.index else 0 for p in par["Parasito"]]
    novos = tabela_parasitos_padrao(contagem[~contagem["Parasito"].isin(par["Parasito"])])
    return pd.concat([par, novos], ignore_index=True)[PAR_COLS]


# ----------------------------------------------------------------------
# tabelas editadas -> AnalysisConfig
# ----------------------------------------------------------------------
def construir_config(parasitos: pd.DataFrame, metodos: pd.DataFrame, n_amostras, criterios=None) -> tuple[AnalysisConfig, list[str]]:
    avisos: list[str] = []

    met_tuplas = []
    nomes_usados = set()
    for _, r in metodos.iterrows():
        if not _as_bool(r["Incluir"]):
            continue
        col = _texto(r["Coluna"])
        nome = _texto(r["Método"]) or _nome_padrao_metodo(col)
        base, i = nome, 2
        while nome in nomes_usados:
            nome = f"{base} ({i})"
            i += 1
        nomes_usados.add(nome)
        dom = _dominio(_as_amostra(r["Amostra"]))
        status = "status_lamina" if dom == DOMINIO_LAMINA else "status_amostra"
        met_tuplas.append((col, nome, status, dom))

    excluidos, renomear, categorias = set(), {}, {}
    por_alvo: dict[str, list[tuple[str, str]]] = {}
    for _, r in parasitos.iterrows():
        p = _texto(r["Parasito"])
        if not p:
            continue
        if not _as_bool(r["Incluir"]):
            excluidos.add(p)
            continue
        alvo = _texto(r.get("Agrupar como", "")) or p
        if alvo != p:
            renomear[p] = alvo
        cat = r["Classificação"] if r["Classificação"] in CATEGORIAS_VALIDAS else _as_categoria(r["Classificação"])
        por_alvo.setdefault(alvo, []).append((p, cat))
    for alvo, itens in por_alvo.items():
        # vale a classificação da linha cujo nome é o próprio nome final; senão, a 1ª
        escolhido = next(((p, c) for p, c in itens if p == alvo), itens[0])
        categorias[alvo] = escolhido[1]
        if len({c for _, c in itens}) > 1:
            avisos.append(f"'{alvo}': nomes agrupados com classificações diferentes — "
                          f"vale '{escolhido[1]}' (da linha '{escolhido[0]}').")

    nao_class = sorted(e for e, c in categorias.items() if c == "Não classificado")
    if nao_class:
        txt = ", ".join(nao_class)
        avisos.append("Sem classificação (não entram nos gráficos de patogênicos): " + txt + ("" if txt.endswith(".") else "."))

    cfg = AnalysisConfig(
        renomear=renomear, excluidos=excluidos, categorias=categorias,
        metodos=met_tuplas, n_amostras=int(n_amostras) if n_amostras else None,
        criterios={**CRITERIOS_PADRAO, **(criterios or {})},
    )
    return cfg, avisos


# ----------------------------------------------------------------------
# exportação
# ----------------------------------------------------------------------
def _abas_config(parasitos: pd.DataFrame, metodos: pd.DataFrame, n_amostras, criterios=None) -> dict[str, pd.DataFrame]:
    par = parasitos.copy()
    par["Incluir"] = par["Incluir"].apply(lambda b: "Sim" if _as_bool(b) else "Não")
    par = par[[c for c in ["Incluir", "Parasito", "Classificação", "Agrupar como"] if c in par.columns]]
    met = metodos.copy()
    met["Incluir"] = met["Incluir"].apply(lambda b: "Sim" if _as_bool(b) else "Não")
    met = met[MET_COLS]
    amo = pd.DataFrame(
        [[PARAM_N_AMOSTRAS, int(n_amostras) if n_amostras else "",
          "Coletas P1..Pn consideradas por paciente. Coletas acima desse número são ignoradas."]],
        columns=["Parâmetro", "Valor", "Observação"],
    )
    c = {**CRITERIOS_PADRAO, **(criterios or {})}
    cri = pd.DataFrame(
        [[rot, _criterio_valor_planilha(ch, c[ch]), expl] for ch, rot, expl in CRITERIOS_ROTULOS],
        columns=["Critério de inclusão", "Valor", "Observação"],
    )
    return {SHEET_PARASITOS: par, SHEET_METODOS: met, SHEET_AMOSTRAS: amo, SHEET_CRITERIOS: cri}


def _ajustar_larguras(ws):
    for col in ws.columns:
        largura = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, largura + 2), 60)


def config_xlsx_bytes(parasitos, metodos, n_amostras, criterios=None) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nome, d in _abas_config(parasitos, metodos, n_amostras, criterios).items():
            d.to_excel(w, sheet_name=nome, index=False)
            _ajustar_larguras(w.sheets[nome])
    return estilizar_config(buf.getvalue(), CATEGORIAS_VALIDAS, AMOSTRAS_VALIDAS)


def planilha_com_config_bytes(original: bytes, parasitos, metodos, n_amostras, criterios=None) -> bytes:
    """Planilha original + abas Config_* (substitui as antigas, se houver)."""
    abas = _abas_config(parasitos, metodos, n_amostras, criterios)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(original))
        for s in list(wb.sheetnames):
            if _key(s) in {_key(x) for x in CONFIG_SHEETS}:
                del wb[s]
        for nome, d in abas.items():
            ws = wb.create_sheet(nome)
            ws.append(list(d.columns))
            for row in d.itertuples(index=False):
                ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
            _ajustar_larguras(ws)
        buf = io.BytesIO()
        wb.save(buf)
        return estilizar_config(buf.getvalue(), CATEGORIAS_VALIDAS, AMOSTRAS_VALIDAS)
    except Exception:  # noqa: BLE001 — ex.: .xls antigo; reescreve só os dados
        xls = pd.ExcelFile(io.BytesIO(original))
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            for s in xls.sheet_names:
                if _key(s) in {_key(x) for x in CONFIG_SHEETS}:
                    continue
                pd.read_excel(xls, sheet_name=s).to_excel(w, sheet_name=s[:31], index=False)
            for nome, d in abas.items():
                d.to_excel(w, sheet_name=nome, index=False)
        return estilizar_config(buf.getvalue(), CATEGORIAS_VALIDAS, AMOSTRAS_VALIDAS)


def abas_config_modelo() -> dict[str, pd.DataFrame]:
    """Abas de configuração do MODELO de planilha: todos os parasitos e métodos
    que o sistema já conhece, com a classificação padrão (editável)."""
    especies = sorted(set(PARASITE_MAP.values()))
    par = pd.DataFrame(
        [{"Incluir": True, "Parasito": e, "Classificação": categoria_de(e), "Agrupar como": ""} for e in especies]
    )
    met = pd.DataFrame(
        [{"Incluir": True, "Coluna": col, "Método": nome, "Amostra": _amostra(dom)}
         for col, nome, _, dom in METHOD_CATALOG],
        columns=MET_COLS,
    )
    return _abas_config(par, met, 3)
