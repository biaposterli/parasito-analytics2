"""
Acabamento visual da planilha-modelo — identidade Pirajá.

Recebe o .xlsx gerado pelo pandas (dados e posições de células já definidos) e
aplica só formatação: cores, fontes, larguras, painéis congelados e listas
suspensas. NÃO muda nenhum cabeçalho nem a posição dos dados que o sistema lê
(abas Dados e Config_*), então a leitura da planilha continua idêntica.
A única aba reorganizada é a "Leia-me", que é só texto e não é lida pelo sistema.
"""
import io
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Paleta "Mata, barro e papel"
MATA = "11483D"
FOLHA = "328567"
COBRE = "9C4A2F"
PAPEL = "F4F1EA"
SUPERFICIE = "FBF9F4"
TINTA = "1B2421"
TINTA_SUAVE = "4E5B55"
TINTA_CLARA = "7A857F"
LINHA = "DCD6C8"
FOLHA_TINTA = "E3EEE8"
COBRE_TINTA = "F3E4DB"

# Fontes que existem em qualquer Excel/LibreOffice (a planilha não embute
# fontes): Georgia faz o papel da Fraunces nos títulos, Calibri o do texto.
SERIFA = "Georgia"
SANS = "Calibri"

_fino = Side(style="thin", color=LINHA)
BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
BORDA_BAIXO = Border(bottom=_fino)


def _fill(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


def _cabecalho(ws, linha=1, cor=MATA, altura=34, colunas=None):
    for cell in ws[linha]:
        if cell.value is None:
            continue
        if colunas and cell.column not in colunas:
            continue
        cell.fill = _fill(cor)
        cell.font = Font(name=SANS, bold=True, color=PAPEL, size=11)
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
        cell.border = BORDA
    ws.row_dimensions[linha].height = altura


def _corpo(ws, primeira_linha=2, zebra=True, wrap=True):
    for r, row in enumerate(ws.iter_rows(min_row=primeira_linha), start=0):
        for cell in row:
            cell.font = Font(name=SANS, size=11, color=TINTA)
            cell.alignment = Alignment(vertical="top", wrap_text=wrap, indent=1)
            cell.border = BORDA
            cell.fill = _fill(SUPERFICIE if (zebra and r % 2) else "FFFFFF")


def _larguras(ws, larguras: dict):
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w


def _lista(ws, valores, intervalo, titulo, msg):
    """Lista suspensa que SUGERE valores mas aceita outros (aviso, não bloqueio)."""
    dv = DataValidation(
        type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True,
        showErrorMessage=True, errorStyle="warning",
        errorTitle=titulo, error=msg, showInputMessage=False,
    )
    ws.add_data_validation(dv)
    dv.add(intervalo)


def _leia_me(ws, logo_path):
    titulo_coluna = ws["A1"].value or "Leia antes de preencher"
    texto = ws["A2"].value or ""
    ws.delete_rows(1, ws.max_row)

    ws.sheet_view.showGridLines = False
    _larguras(ws, {"A": 3, "B": 110})
    linha = 1
    if logo_path and Path(logo_path).exists():
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(str(logo_path))
            img.height, img.width = 62, int(62 * img.width / img.height)
            ws.add_image(img, "B2")
            ws.row_dimensions[2].height = 50
            linha = 5
        except Exception:  # noqa: BLE001
            linha = 2
    else:
        linha = 2

    ws.cell(linha, 2, "Modelo de planilha — levantamento de parasitoses intestinais").font = Font(
        name=SERIFA, size=20, bold=True, color=MATA)
    ws.row_dimensions[linha].height = 30
    ws.cell(linha + 1, 2, titulo_coluna.upper()).font = Font(name=SANS, size=10, bold=True, color=COBRE)
    ws.cell(linha + 1, 2).border = BORDA_BAIXO
    linha += 3

    for par in [p.strip() for p in str(texto).split("\n\n") if p.strip()]:
        c = ws.cell(linha, 2, par)
        destaque = par[:40].isupper() or par.startswith("CONFIGURAÇÃO")
        c.font = Font(name=SANS, size=11.5, color=TINTA if not destaque else MATA, bold=False)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if destaque:
            c.fill = _fill(FOLHA_TINTA)
        # ~105 caracteres por linha na largura 110
        ws.row_dimensions[linha].height = max(18, 16 * (len(par) // 105 + 1) + 6)
        linha += 2

    ws.cell(linha, 2, "Abas deste arquivo").font = Font(name=SERIFA, size=14, bold=True, color=MATA)
    linha += 1
    abas = [
        ("Dados", "onde você registra os resultados — uma linha por coleta (P1, P2, P3…) de cada paciente."),
        ("Legenda", "o que é cada coluna e os valores aceitos."),
        ("Especies_Reconhecidas", "parasitos que o sistema já reconhece e as grafias aceitas."),
        ("Config_*", "configuração da análise (opcional): parasitos, métodos e nº de amostras."),
    ]
    for nome, desc in abas:
        c = ws.cell(linha, 2, f"{nome}  —  {desc}")
        c.font = Font(name=SANS, size=11, color=TINTA_SUAVE)
        linha += 1
    ws.sheet_properties.tabColor = MATA


def _dados(ws, colunas_lamina=()):
    cab = [c.value for c in ws[1]]
    n = len(cab)
    for i, nome in enumerate(cab, start=1):
        if nome is None:
            continue
        nome = str(nome)
        if nome in colunas_lamina:
            cor = COBRE   # métodos de lâmina
        elif nome.startswith("metodo_"):
            cor = FOLHA   # métodos de fezes
        else:
            cor = MATA
        _cabecalho(ws, 1, cor=cor, altura=36, colunas={i})
        largura = 30 if nome.startswith("metodo_") else {"id_paciente": 14, "coleta": 10,
                                                         "nome_paciente": 24, "nome_responsavel": 24,
                                                         "status_amostra": 17, "status_lamina": 17,
                                                         "observacoes": 46}.get(nome, 18)
        ws.column_dimensions[get_column_letter(i)].width = largura

    # linhas de exemplo: cinza-claro e itálico, para ficar claro que são só referência
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=SANS, size=11, italic=True, color=TINTA_CLARA)
            cell.fill = _fill(PAPEL)
            cell.border = BORDA
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws["A2"].comment = Comment(
        "Linhas de exemplo (F-001 e F-002): servem só de referência. Apague-as antes de "
        "enviar, senão entram na análise.", "Pirajá")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n)}{max(ws.max_row, 2)}"

    col = {str(c.value): c.column_letter for c in ws[1] if c.value}
    for campo in ("status_amostra", "status_lamina"):
        if campo in col:
            _lista(ws, ["Entregue", "Não entregue"], f"{col[campo]}2:{col[campo]}2000",
                   "Status", "Valores esperados: Entregue ou Não entregue.")
    if "coleta" in col:
        _lista(ws, [f"P{i}" for i in range(1, 7)], f"{col['coleta']}2:{col['coleta']}2000",
               "Coleta", "Use P1, P2, P3… para identificar a coleta.")
    ws.sheet_properties.tabColor = COBRE


def _legenda(ws):
    _cabecalho(ws)
    _corpo(ws)
    _larguras(ws, {"A": 26, "B": 62, "C": 80})
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(name=SANS, size=11, bold=True, color=MATA)
    ws.freeze_panes = "A2"


def _especies(ws):
    # A1 = título da nota, A2 = nota; tabela a partir da linha 4 (cabeçalho)
    ws.merge_cells("A2:C2")
    ws["A1"].font = Font(name=SERIFA, size=14, bold=True, color=MATA)
    ws["A2"].font = Font(name=SANS, size=11, color=TINTA_SUAVE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A2"].fill = _fill(FOLHA_TINTA)
    ws.row_dimensions[2].height = 125
    _cabecalho(ws, linha=4)
    _corpo(ws, primeira_linha=5)
    _larguras(ws, {"A": 38, "B": 20, "C": 70})
    for row in ws.iter_rows(min_row=5):
        row[0].font = Font(name=SANS, size=11, italic=True, color=TINTA)
        cat = str(row[1].value or "")
        if cat.startswith("Patog"):
            row[1].fill, row[1].font = _fill(COBRE_TINTA), Font(name=SANS, size=11, bold=True, color=COBRE)
        elif cat.startswith("Comens"):
            row[1].fill, row[1].font = _fill(FOLHA_TINTA), Font(name=SANS, size=11, bold=True, color=FOLHA)
    ws.freeze_panes = "A5"


def _config(ws, nome, categorias, amostras):
    _cabecalho(ws, cor=FOLHA)
    _corpo(ws, wrap=False)
    for col in ws.columns:
        largura = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, largura + 4), 70)
    ws.freeze_panes = "A2"
    col = {str(c.value): c.column_letter for c in ws[1] if c.value}
    ultima = max(ws.max_row, 2) + 200
    if "Incluir" in col:
        _lista(ws, ["Sim", "Não"], f"{col['Incluir']}2:{col['Incluir']}{ultima}",
               "Incluir", "Use Sim ou Não.")
    if "Classificação" in col and categorias:
        _lista(ws, list(categorias), f"{col['Classificação']}2:{col['Classificação']}{ultima}",
               "Classificação", "Valores esperados: " + ", ".join(categorias) + ".")
    if "Amostra" in col and amostras:
        _lista(ws, list(amostras), f"{col['Amostra']}2:{col['Amostra']}{ultima}",
               "Amostra", "Valores esperados: " + ", ".join(amostras) + ".")
    if nome == "Config_Amostras":
        for row in ws.iter_rows(min_row=2):
            row[0].font = Font(name=SANS, size=11, bold=True, color=MATA)
            if len(row) > 1:
                row[1].font = Font(name=SERIFA, size=13, bold=True, color=COBRE)
                row[1].alignment = Alignment(horizontal="center", vertical="top")
    ws.sheet_properties.tabColor = FOLHA


def estilizar_modelo(xlsx: bytes, logo_path=None, categorias=(), amostras=(), colunas_lamina=()) -> bytes:
    """Aplica a identidade visual à planilha-modelo. Em caso de qualquer erro,
    devolve o arquivo original sem formatação (nunca quebra o download)."""
    try:
        wb = load_workbook(io.BytesIO(xlsx))
        for ws in wb.worksheets:
            ws.sheet_view.zoomScale = 110
            # impressão: paisagem, ajustada à largura da página
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            if ws.title == "Leia-me":
                _leia_me(ws, logo_path)
            elif ws.title == "Dados":
                _dados(ws, colunas_lamina)
            elif ws.title == "Legenda":
                _legenda(ws)
            elif ws.title == "Especies_Reconhecidas":
                _especies(ws)
            elif ws.title.startswith("Config_"):
                _config(ws, ws.title, categorias, amostras)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except Exception:  # noqa: BLE001
        return xlsx
